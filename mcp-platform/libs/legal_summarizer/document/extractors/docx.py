"""Извлечение из DOCX.

В отличие от PDF, DOCX хранит структуру в самом файле: абзац — это элемент
``w:p``, таблица — ``w:tbl``. Поэтому извлекатель отдаёт единицу на элемент и
почти ничего не выдумывает; разделение получает уже готовые абзацы и, как
правило, ничего в них не меняет.

Обход ``body.iterchildren()`` даёт честный document order (абзацы и таблицы
вперемешку), а не «сначала все абзацы, потом все таблицы». python-docx
хранит ``doc.paragraphs`` и ``doc.tables`` раздельно и взаимный порядок не
помнит, поэтому за позициями в обоих списках идём по XML-элементам напрямую.

``sectPr`` пропускаем — это не контент.
"""

from __future__ import annotations

from pathlib import Path

from libs.legal_summarizer.document.extraction import (
    KIND_TABLE,
    ExtractionResult,
    RawUnit,
    render_table,
)

__all__ = ["DocxExtractor"]


class DocxExtractor:
    """DOCX → сырые единицы (``w:p`` и ``w:tbl`` в document order)."""

    format_name = "docx"

    def extract(self, path: Path) -> ExtractionResult:
        from docx import Document

        doc = Document(str(path))
        para_by_xml = {id(p._p): (i, p) for i, p in enumerate(doc.paragraphs)}
        table_by_xml = {id(t._tbl): (i, t) for i, t in enumerate(doc.tables)}

        units: list[RawUnit] = []
        for child in doc.element.body.iterchildren():
            local_tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

            if local_tag == "p":
                entry = para_by_xml.get(id(child))
                if entry is None:
                    continue
                paragraph_index, para = entry
                text = (para.text or "").strip()
                if not text:
                    continue
                units.append(
                    RawUnit(
                        text=text,
                        paragraph_index=paragraph_index,
                        block_metadata={
                            "style": para.style.name if para.style else ""
                        },
                    )
                )

            elif local_tag == "tbl":
                entry = table_by_xml.get(id(child))
                if entry is None:
                    continue
                table_index, table = entry
                rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                table_text = render_table(rows)
                if not table_text.strip():
                    continue
                units.append(
                    RawUnit(
                        text=table_text,
                        kind=KIND_TABLE,
                        table_index=table_index,
                        block_metadata={"row_count": len(rows)},
                    )
                )

        return ExtractionResult(
            units=tuple(units),
            page_count=1,
            title=self._title(doc),
        )

    @staticmethod
    def _title(doc) -> str | None:
        try:
            title = doc.core_properties.title
        except Exception:
            return None
        if not title:
            return None
        return str(title).strip() or None