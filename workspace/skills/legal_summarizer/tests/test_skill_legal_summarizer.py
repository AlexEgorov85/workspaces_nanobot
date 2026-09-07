"""Behavioral smoke tests для legal_summarizer (Phase 2B — Skill).

Перенесено из tests/test_skill_legal_summarizer.py в рамках миграции
Skill к целевой структуре (runtime в scripts/).

Содержит только тесты, которые проходят с текущим API:
* output/prepare_output (pure)
* skill_config defaults
* load_text (DocumentLoader)
* DocumentLoader (DOCX)
* strip_think_blocks (pure)
* SKILL.md contract
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import pytest

_SKILL_ROOT = Path(__file__).resolve().parents[1]
_SCRIPTS_DIR = _SKILL_ROOT / "scripts"
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))
_PROJECT_ROOT = _SKILL_ROOT.parents[2]
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))


# ---------------------------------------------------------------------------
# summarizer.run() — короткий документ (пустой текст → EMPTY_DOCUMENT)
# ---------------------------------------------------------------------------


def test_run_empty_text_returns_failed(tmp_path):
    """Пустой текст → status=failed с code=EMPTY_DOCUMENT."""
    from application.service import run
    p = tmp_path / "empty.txt"
    p.write_text("   \n\n  \n", encoding="utf-8")
    result = run(
        "", length="brief", workspace_root=tmp_path,
        document_path=str(p),
    )
    assert result["status"] == "failed"
    assert result["error"]["code"] == "EMPTY_DOCUMENT"


# ---------------------------------------------------------------------------
# load_text / load_structure
# ---------------------------------------------------------------------------


def test_load_text_txt_success(tmp_path):
    """load_text для .txt читает содержимое."""
    from application.service import load_text
    p = tmp_path / "contract.txt"
    p.write_text(
        "Договор аренды.\n\nАрендодатель сдаёт помещение.\n",
        encoding="utf-8",
    )
    text = load_text(p)
    assert "Договор аренды" in text


def test_load_text_missing_raises(tmp_path):
    """load_text для несуществующего файла → FileNotFoundError."""
    import pytest
    from application.service import load_text
    with pytest.raises(FileNotFoundError):
        load_text(tmp_path / "missing.pdf")


def test_load_text_unknown_extension_raises(tmp_path):
    """load_text для неизвестного расширения → ValueError."""
    import pytest
    from application.service import load_text
    p = tmp_path / "data.bin"
    p.write_bytes(b"\x00\x01\x02" * 16)
    with pytest.raises(ValueError):
        load_text(p)


def test_load_text_brief_mode_for_txt(tmp_path):
    """mode='brief' для .txt = mode='full'."""
    from application.service import load_text
    p = tmp_path / "contract.txt"
    p.write_text(
        "Договор аренды.\n\nПункт 1.\n\nПункт 2.",
        encoding="utf-8",
    )
    text = load_text(p, mode="brief")
    assert "Договор аренды" in text
    text_full = load_text(p, mode="full")
    assert text == text_full


def test_load_structure_returns_title_and_text(tmp_path):
    """DocumentLoader для DOCX возвращает title + blocks."""
    from docx import Document
    from document.loader import DocumentLoader

    p = tmp_path / "contract.docx"
    d = Document()
    d.core_properties.title = "Договор поставки №7"
    d.add_paragraph("1. Поставщик передаёт товар покупателю.")
    d.add_paragraph("2. Покупатель оплачивает товар в течение 5 дней.")
    d.save(str(p))

    doc = DocumentLoader().load(p)
    assert "Договор поставки №7" == doc.title
    blocks_text = "\n\n".join(b.content for b in doc.blocks)
    assert "Поставщик" in blocks_text


# ---------------------------------------------------------------------------
# output / prepare_output
# ---------------------------------------------------------------------------


def test_prepare_output_completed():
    """prepare_output для completed выдаёт subject/summary/chunks."""
    from output.presenter import prepare_output
    result = {
        "status": "completed",
        "operation_id": "op_test_001",
        "result": {
            "subject": "Договор аренды.",
            "summary": "...",
            "length": "medium",
            "chars_in": 100,
            "chunks": 3,
            "context_batches": 2,
            "sections": 5,
            "strategy": "map_reduce_hierarchical",
        },
        "stats": {"actual_llm_calls": 4},
    }
    out = prepare_output(result)
    assert out["mode"] == "summarize"
    assert out["status"] == "completed"
    assert out["operation_id"] == "op_test_001"
    assert out["subject"] == "Договор аренды."
    assert out["chunks"] == 3
    assert out["context_batches"] == 2
    assert out["sections"] == 5
    assert out["strategy"] == "map_reduce_hierarchical"


def test_prepare_output_failed():
    """prepare_output для failed пробрасывает error/code."""
    from output.presenter import prepare_output
    result = {
        "status": "failed",
        "operation_id": "op_test_002",
        "error": {
            "code": "EMPTY_DOCUMENT",
            "message": "Документ не содержит текста",
        },
    }
    out = prepare_output(result)
    assert out["status"] == "failed"
    assert out["error"]["code"] == "EMPTY_DOCUMENT"


def test_prepare_output_confirmation_required():
    """prepare_output для confirmation_required: компактный payload,
    без технических чисел.
    """
    from output.presenter import prepare_output
    result = {
        "status": "confirmation_required",
        "operation_id": "op_test_003",
        "summary": {"title": "Договор аренды"},
        "estimate": {
            "min_seconds": 320,
            "max_seconds": 480,
        },
    }
    out = prepare_output(result)
    assert out["status"] == "confirmation_required"
    assert "options" in out
    assert set(out["options"].keys()) == {"brief", "detailed"}
    assert out["options"]["brief"]["min_sec"] < out["options"]["detailed"]["min_sec"]
    assert out["supports_question"] is True
    out_str = str(out)
    assert "chunks_total" not in out_str
    assert "context_batches_total" not in out_str
    assert "estimated_llm_calls" not in out_str
    serialized_len = len(json.dumps(out, ensure_ascii=False))
    assert serialized_len < 400


def test_sanitize_handles_datetime():
    """_sanitize_value сериализует datetime в ISO format."""
    from datetime import datetime
    from output.presenter import _sanitize_value
    out = _sanitize_value({"d": datetime(2024, 1, 15, 10, 30)})
    assert out["d"] == "2024-01-15T10:30:00"


# ---------------------------------------------------------------------------
# skill_config
# ---------------------------------------------------------------------------


def test_skill_config_chunking_defaults_match_project_json():
    """get_chunking_config возвращает defaults из project.json."""
    import llm.config as skill_config
    cfg = skill_config.get_chunking_config()
    assert cfg["chunk_size"] == 100000
    assert cfg["chunk_overlap"] == 0
    assert cfg["single_call_threshold"] == 20000
    assert cfg["chunk_size_input_ratio"] == 0.5


def test_skill_config_cli_matches_project_json():
    """get_cli_config возвращает defaults из project.json."""
    import llm.config as skill_config
    cli = skill_config.get_cli_config()
    assert cli["max_retries"] == 3
    assert cli["timeout_sec"] == 120
    assert skill_config.get_default_length() == "medium"


# ---------------------------------------------------------------------------
# SKILL.md contract
# ---------------------------------------------------------------------------


class TestSkillMarkdownContract:
    @pytest.fixture(scope="class")
    def skill_text(self) -> str:
        path = _SKILL_ROOT / "SKILL.md"
        return path.read_text(encoding="utf-8")

    def test_cli_invocation_in_first_lines(self, skill_text: str):
        head = "\n".join(skill_text.splitlines()[:30])
        assert "cli.py" in head
        assert "--file" in head

    def test_summarize_is_not_a_summary(self, skill_text: str):
        lower = skill_text.lower()
        assert "summarize" in lower
        assert ("не саммари" in lower or "не делает саммари" in lower or
                "не делает llm" in lower)

    def test_mentions_confirm_protocol(self, skill_text: str):
        assert "confirmation_required" in skill_text
        assert "--confirm" in skill_text

    def test_mentions_focus_argument(self, skill_text: str):
        assert "--focus" in skill_text

    def test_forbidden_summarize_direct_call(self, skill_text: str):
        assert (
            "office_files.extract_metadata" in skill_text
            or "office_files.summarize" in skill_text
        )
        assert "❌" in skill_text

    def test_description_mentions_cli(self, skill_text: str):
        assert skill_text.startswith("---")
        end = skill_text.find("\n---\n", 4)
        assert end > 0
        front = skill_text[4:end]
        assert "cli.py" in front
        assert "--file" in front


# ---------------------------------------------------------------------------
# strip_think_blocks (Phase 2B hardening)
# ---------------------------------------------------------------------------


def test_strip_think_blocks_removes_cot():
    """``<think>...</think>`` вырезаются из текста."""
    from llm.sanitize import strip_think_blocks
    text = "<think>\nЭто рассуждение, которое не нужно.\n</think>\nЭто саммари."
    cleaned = strip_think_blocks(text)
    assert "<think>" not in cleaned
    assert "рассуждение" not in cleaned
    assert cleaned == "Это саммари."


def test_strip_think_blocks_no_think_returns_unchanged():
    from llm.sanitize import strip_think_blocks
    text = "Обычное саммари без CoT."
    assert strip_think_blocks(text) == text


def test_strip_think_blocks_multiple():
    from llm.sanitize import strip_think_blocks
    text = "<think>A</think>Полезно.<think>B</think>Конец."
    assert strip_think_blocks(text) == "Полезно.Конец."


def test_strip_think_blocks_unclosed_drops_to_blank_line():
    """Незакрытый ``<think>`` отрезается до первого абзацного разрыва."""
    from llm.sanitize import strip_think_blocks
    text = "<think>\nвнутреннее рассуждение модели\n\nЭто итоговый ответ."
    cleaned = strip_think_blocks(text)
    assert "<think>" not in cleaned
    assert "рассуждение" not in cleaned
    assert cleaned == "Это итоговый ответ."


def test_strip_think_blocks_unclosed_no_blank_drops_all():
    """Незакрытый ``<think>`` без пустой строки → пустая строка."""
    from llm.sanitize import strip_think_blocks
    text = "<think>\nтолько рассуждение без ответа"
    assert strip_think_blocks(text) == ""


def test_strip_think_blocks_mixed_closed_and_unclosed():
    """Смешанный случай: закрытый + незакрытый блоки."""
    from llm.sanitize import strip_think_blocks
    text = "<think>закрытое рассуждение</think>вступление<think>открытое\n\nответ"
    cleaned = strip_think_blocks(text)
    assert "рассуждение" not in cleaned
    assert "<think>" not in cleaned
    assert cleaned == "вступлениеответ"


# ---------------------------------------------------------------------------
# prepare_output partial / cache_stats (pure)
# ---------------------------------------------------------------------------


def test_prepare_output_partial_exposes_failed_batches():
    """``status=partial`` пробрасывает ``partial``/``failed_batches``/``hint``."""
    from output.presenter import prepare_output
    result = {
        "status": "partial",
        "operation_id": "op_x",
        "result": {
            "subject": "Субъект",
            "summary": "Текст.",
            "length": "medium",
            "chars_in": 1000,
            "chunks": 5,
            "context_batches": 3,
            "sections": 0,
            "strategy": "map_reduce_flat",
            "partial": True,
        },
        "stats": {
            "chars_in": 1000,
            "chunks_total": 5,
            "context_batches_total": 3,
            "map_calls": 2,
            "failed_batches": ["cb_000"],
            "partial": True,
        },
    }
    out = prepare_output(result)
    assert out["status"] == "partial"
    assert out["partial"] is True
    assert out["failed_batches"] == ["cb_000"]
    assert "hint" in out
    assert "cb_000" in out["hint"]


def test_prepare_output_includes_cache_stats_for_completed():
    """prepare_output пробрасывает cache_stats в output для completed."""
    from output.presenter import prepare_output
    result = {
        "status": "completed",
        "operation_id": "op_test",
        "result": {
            "subject": "S",
            "summary": "sum",
            "length": "brief",
            "chars_in": 100,
            "chunks": 1,
            "context_batches": 0,
            "sections": 0,
            "strategy": "single",
        },
        "stats": {},
        "cache_stats": {
            "document_id": "abc123",
            "chunks_from_cache": 0,
            "chunks_processed": 1,
            "cache_enabled": False,
        },
    }
    out = prepare_output(result)
    assert "cache_stats" in out
    assert out["cache_stats"]["document_id"] == "abc123"


def test_prepare_output_omits_cache_stats_when_absent():
    from output.presenter import prepare_output
    result = {
        "status": "completed",
        "operation_id": "op_test",
        "result": {
            "subject": "S",
            "summary": "sum",
            "length": "brief",
            "chars_in": 100,
            "chunks": 1,
            "context_batches": 0,
            "sections": 0,
            "strategy": "single",
        },
        "stats": {},
    }
    out = prepare_output(result)
    assert "cache_stats" not in out
