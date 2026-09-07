"""E2E question pipeline: retrieval → selected chunks → LLM input.

Полный путь ``run(question=...)``: LLM-вход должен содержать ТОЛЬКО
выбранные retrieval'ом чанки (или результат lexical/document-head
fallback'а), а не весь документ.

Документ — DOCX (heading-paragraph + body-paragraph на секцию):
canonical pipeline строит секции и по одному чанку на секцию.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

_SKILL_ROOT = Path(__file__).resolve().parents[1]
_SCRIPTS_DIR = _SKILL_ROOT / "scripts"
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))

import docx  # noqa: E402

_CHUNK_ID_RE = re.compile(r"Chunk (\d{3})")


def _build_docx(
    tmp_path: Path,
    name: str,
    sections: int,
    marker_in: frozenset[int] = frozenset(),
    marker_word: str = "контрагентство",
) -> tuple[Path, str]:
    """N numbered sections; marker_word — только в телах marker_in."""
    doc = docx.Document()
    for i in range(1, sections + 1):
        doc.add_paragraph(f"{i}. Раздел {i}")
        body = "Общие положения и порядок действий раздела. "
        if i in marker_in:
            body += f"Специальное правило: {marker_word} применений. "
        doc.add_paragraph(body * 20)
    p = tmp_path / name
    doc.save(str(p))
    import legal_summarizer.application.service as summarizer
    text = summarizer.load_text(p)
    return p, text


def _install_recording_llm(monkeypatch, recorded: dict) -> None:
    """Mock LLM-вызовов, собирающие chunk-IDs, реально ушедшие в LLM."""
    import legal_summarizer.llm.calls as llm_calls
    import legal_summarizer.execution.pipeline as _pipeline_mod
    import legal_summarizer.application.service as summarizer
    def _fake_batch(chunks, *, chunks_total, structure, length, question=None):
        recorded["ids"].update(c.chunk_id for c in chunks)
        return {c.chunk_id: f"summary {c.chunk_id}" for c in chunks}

    def _fake_section(path, heading, text, *, length, question=None):
        recorded["ids"].update(_CHUNK_ID_RE.findall(text))
        return "section summary"

    def _fake_doc(text, *, length, focus, structure, question=None):
        recorded["ids"].update(_CHUNK_ID_RE.findall(text))
        return "doc summary"

    monkeypatch.setattr(llm_calls, "llm_batch", _fake_batch)
    monkeypatch.setattr(llm_calls, "llm_section_reduce", _fake_section)
    monkeypatch.setattr(llm_calls, "llm_document_reduce", _fake_doc)
    monkeypatch.setattr(summarizer, "_llm_batch", _fake_batch)
    monkeypatch.setattr(summarizer, "_llm_section_reduce", _fake_section)
    monkeypatch.setattr(summarizer, "_llm_document_reduce", _fake_doc)
    monkeypatch.setattr(_pipeline_mod, "_llm_batch", _fake_batch)


def test_question_llm_input_only_retrieved_chunks(tmp_path, monkeypatch):
    """Маркер в 2 из 8 секций: LLM видит только эти 2 чанка."""
    import legal_summarizer.application.service as summarizer
    recorded = {"ids": set()}
    _install_recording_llm(monkeypatch, recorded)

    p, text = _build_docx(tmp_path, "q1.docx", 8, marker_in=frozenset({3, 7}))

    insp = summarizer.inspect(text, document_path=str(p))
    expected = {c.chunk_id for c in insp.chunks if "контрагентство" in c.text}
    assert len(expected) == 2, f"setup: нужно ровно 2 marker-чанка, {expected}"
    # Sanity: retrieval-индекс находит терм по точному совпадению.
    assert [h.chunk_id for h in insp.analysis.retrieve("контрагентство")] and \
        {h.chunk_id for h in insp.analysis.retrieve("контрагентство")} == expected

    result = summarizer.run(
        text, question="контрагентство",
        document_path=str(p), workspace_root=tmp_path, confirmed=True,
    )
    assert result["status"] == "completed", result
    assert set(recorded["ids"]) == expected, (
        f"LLM увидел чанки {sorted(recorded['ids'])}, ожидаются {sorted(expected)}"
    )
    assert result["result"]["chunks"] == len(expected)


def test_question_respects_max_chunks_per_question(tmp_path, monkeypatch):
    """max_chunks_per_question=1 → ровно один (top) чанк уходит в LLM."""
    import legal_summarizer.application.service as summarizer
    monkeypatch.setattr(summarizer, "_resolve_max_chunks", lambda: 1)
    recorded = {"ids": set()}
    _install_recording_llm(monkeypatch, recorded)

    p, text = _build_docx(tmp_path, "q2.docx", 8, marker_in=frozenset({3, 7}))

    insp = summarizer.inspect(text, document_path=str(p))
    marker_ids = sorted(c.chunk_id for c in insp.chunks if "контрагентство" in c.text)
    assert len(marker_ids) == 2

    result = summarizer.run(
        text, question="контрагентство",
        document_path=str(p), workspace_root=tmp_path, confirmed=True,
    )
    assert result["status"] == "completed", result
    # Тира по score решается chunk_id'ом → топ = минимальный id.
    assert set(recorded["ids"]) == {marker_ids[0]}, (
        f"expected top-1 {marker_ids[0]}, got {sorted(recorded['ids'])}"
    )


def test_question_lexical_fallback_when_retrieval_empty(tmp_path, monkeypatch):
    """Inverted index пуст → relaxed lexical (4-буквенный префикс)
    выбирает чанки по подстроке."""
    import legal_summarizer.application.service as summarizer
    recorded = {"ids": set()}
    _install_recording_llm(monkeypatch, recorded)

    # В тексте — форма "контрагентские"; запрос "контрагентский" не входит
    # в inverted index как точный терм, поэтому retrieval пуст и
    # срабатывает lexical fallback по префиксу.
    p, text = _build_docx(
        tmp_path, "q3.docx", 8,
        marker_in=frozenset({2, 5}), marker_word="контрагентские",
    )

    insp = summarizer.inspect(text, document_path=str(p))
    expected = {c.chunk_id for c in insp.chunks if "контрагентские" in c.text}
    assert len(expected) == 2
    # Sanity-check: retrieval-индекс не содержит точного терма запроса.
    assert not insp.analysis.retrieve("контрагентский"), (
        "setup: retrieval должен быть пуст"
    )

    result = summarizer.run(
        text, question="контрагентский",
        document_path=str(p), workspace_root=tmp_path, confirmed=True,
    )
    assert result["status"] == "completed", result
    assert set(recorded["ids"]) == expected, (
        f"LLM увидел чанки {sorted(recorded['ids'])}, ожидаются {sorted(expected)}"
    )


def test_question_full_miss_falls_back_to_document_head(tmp_path, monkeypatch):
    """Ни retrieval, ни lexical не нашли ничего → bounded doc-head fallback
    (первые question_fallback_max_chunks чанков по умолчанию)."""
    import legal_summarizer.application.service as summarizer
    recorded = {"ids": set()}
    _install_recording_llm(monkeypatch, recorded)

    p, text = _build_docx(tmp_path, "q4.docx", 8)
    insp = summarizer.inspect(text, document_path=str(p))
    assert len(insp.chunks) == 8

    result = summarizer.run(
        text, question="незнайдённое",
        document_path=str(p), workspace_root=tmp_path, confirmed=True,
    )
    assert result["status"] == "completed", result
    assert set(recorded["ids"]) == {c.chunk_id for c in insp.chunks}


def test_relaxed_lexical_fallback_unit():
    """_relaxed_lexical_fallback: prefix-match, лимит, empty-случаи."""
    import legal_summarizer.application.service as summarizer
    class _Chunk:
        def __init__(self, text: str) -> None:
            self.text = text

    chunks = [
        _Chunk("финансовые сроки исполнения"),
        _Chunk("общие положения"),
        _Chunk("финальные нормы"),
    ]
    res = summarizer._relaxed_lexical_fallback(
        "финальный", chunks, max_chunks=2,
    )
    assert [c.text for c in res] == [
        "финансовые сроки исполнения",
        "финальные нормы",
    ]

    limited = summarizer._relaxed_lexical_fallback(
        "финальный", chunks, max_chunks=1,
    )
    assert len(limited) == 1

    assert summarizer._relaxed_lexical_fallback("zzz", chunks, max_chunks=2) is None
    assert summarizer._relaxed_lexical_fallback("финальный", [], max_chunks=2) is None
    assert summarizer._relaxed_lexical_fallback("финальный", chunks, max_chunks=0) is None
    assert summarizer._relaxed_lexical_fallback("", chunks, max_chunks=2) is None
